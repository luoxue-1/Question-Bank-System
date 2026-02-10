import tkinter as tk
from tkinter import ttk, messagebox
from question_manager import QuestionManager
from paper_generator import PaperGenerator
from exporter import Exporter

class QuizApp:
    """题库系统GUI应用"""
    
    def __init__(self, root):
        """
        初始化应用
        
        Args:
            root: 根窗口
        """
        self.root = root
        self.root.title("个人题库系统")
        self.root.geometry("1000x700")
        
        # 初始化管理器
        self.question_manager = QuestionManager()
        self.paper_generator = PaperGenerator()
        self.exporter = Exporter()
        
        # 保存当前选中的题目ID
        self.current_question_id = None
        # 保存当前生成的试卷和答案
        self.current_paper = None
        self.current_answers = None
        # 连续添加模式标志
        self.continuous_add_mode = False
        
        # 创建标签页
        self.notebook = ttk.Notebook(root)
        self.notebook.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)
        
        # 创建题目管理页
        self.question_frame = ttk.Frame(self.notebook)
        self.notebook.add(self.question_frame, text="题目管理")
        
        # 创建组卷导出页
        self.paper_frame = ttk.Frame(self.notebook)
        self.notebook.add(self.paper_frame, text="组卷导出")
        
        # 初始化题目管理页
        self._init_question_page()
        
        # 初始化组卷导出页
        self._init_paper_page()
        
        # 加载题库列表
        self._load_banks()
        
        # 加载题目数据
        self._load_questions()
    
    def _init_question_page(self):
        """
        初始化题目管理页
        """
        # 顶部筛选区域
        filter_frame = ttk.Frame(self.question_frame)
        filter_frame.pack(fill=tk.X, pady=5)
        
        # 题库选择
        ttk.Label(filter_frame, text="题库:").pack(side=tk.LEFT, padx=5)
        self.bank_var = tk.StringVar()
        self.bank_combobox = ttk.Combobox(filter_frame, textvariable=self.bank_var, width=15)
        self.bank_combobox.pack(side=tk.LEFT, padx=5)
        self.bank_combobox.bind("<<ComboboxSelected>>", self._filter_by_bank)
        
        # 题库管理按钮
        ttk.Button(filter_frame, text="题库管理", command=self._manage_banks).pack(side=tk.LEFT, padx=5)
        
        # 知识点筛选
        ttk.Label(filter_frame, text="知识点筛选:").pack(side=tk.LEFT, padx=5)
        self.topic_var = tk.StringVar()
        self.topic_combobox = ttk.Combobox(filter_frame, textvariable=self.topic_var, width=20)
        self.topic_combobox.pack(side=tk.LEFT, padx=5)
        self.topic_combobox.bind("<<ComboboxSelected>>", self._filter_questions)
        
        # 搜索
        ttk.Label(filter_frame, text="搜索:").pack(side=tk.LEFT, padx=5)
        self.search_var = tk.StringVar()
        search_entry = ttk.Entry(filter_frame, textvariable=self.search_var, width=30)
        search_entry.pack(side=tk.LEFT, padx=5)
        ttk.Button(filter_frame, text="搜索", command=self._search_questions).pack(side=tk.LEFT, padx=5)
        ttk.Button(filter_frame, text="清空", command=self._clear_search).pack(side=tk.LEFT, padx=5)
        
        # 中间列表区域
        list_frame = ttk.Frame(self.question_frame)
        list_frame.pack(fill=tk.BOTH, expand=True, pady=5)
        
        # 创建滚动条
        scrollbar = ttk.Scrollbar(list_frame, orient=tk.VERTICAL)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        
        # 创建题目列表，设置为多选模式
        columns = ("id", "bank", "topic", "content")
        self.question_tree = ttk.Treeview(list_frame, columns=columns, show="headings", yscrollcommand=scrollbar.set, selectmode="extended")
        self.question_tree.pack(fill=tk.BOTH, expand=True)
        scrollbar.config(command=self.question_tree.yview)
        
        # 设置列标题
        self.question_tree.heading("id", text="ID")
        self.question_tree.heading("bank", text="题库")
        self.question_tree.heading("topic", text="知识点")
        self.question_tree.heading("content", text="题目内容")
        
        # 添加排序功能
        self.question_tree.heading("id", text="ID", command=lambda: self._sort_column("id"))
        
        # 设置列宽
        self.question_tree.column("id", width=50)
        self.question_tree.column("bank", width=100)
        self.question_tree.column("topic", width=120)
        self.question_tree.column("content", width=550)
        
        # 绑定选择事件
        self.question_tree.bind("<<TreeviewSelect>>", self._on_question_select)
        
        # 底部编辑区域
        edit_frame = ttk.Frame(self.question_frame)
        edit_frame.pack(fill=tk.BOTH, expand=True, pady=5)
        
        # 左侧编辑区域
        left_frame = ttk.Frame(edit_frame)
        left_frame.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=5)
        
        # 知识点输入区域，包含输入框和退出按钮
        topic_frame = ttk.Frame(left_frame)
        topic_frame.pack(fill=tk.X, pady=2)
        
        ttk.Label(topic_frame, text="知识点:").pack(side=tk.LEFT, padx=5)
        
        self.edit_topic_var = tk.StringVar()
        # 初始化Combobox为可编辑模式
        # 使用'state'参数设置为'normal'，确保完全可编辑
        self.topic_combobox_edit = ttk.Combobox(topic_frame, textvariable=self.edit_topic_var, width=40)
        # 显式设置为可编辑状态
        self.topic_combobox_edit.configure(state='normal')
        # 初始加载所有知识点
        all_topics = self.question_manager.get_all_topics()
        self.topic_combobox_edit['values'] = all_topics
        self.topic_combobox_edit.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=5)
        # 为知识点输入框添加自动完成功能
        self.topic_combobox_edit.bind("<KeyRelease>", self._on_topic_input)
        
        # 添加退出按钮，用于清空知识点输入
        ttk.Button(topic_frame, text="退出", command=self._exit_topic, width=10).pack(side=tk.RIGHT, padx=5)
        
        ttk.Label(left_frame, text="题目内容:").pack(anchor=tk.W, pady=2)
        self.edit_content_var = tk.StringVar()
        content_text = tk.Text(left_frame, height=10, width=50)
        content_text.pack(fill=tk.BOTH, expand=True, pady=2)
        self.edit_content_text = content_text
        
        ttk.Label(left_frame, text="答案:").pack(anchor=tk.W, pady=2)
        self.edit_answer_var = tk.StringVar()
        answer_text = tk.Text(left_frame, height=5, width=50)
        answer_text.pack(fill=tk.BOTH, expand=True, pady=2)
        self.edit_answer_text = answer_text
        
        # 右侧按钮区域
        right_frame = ttk.Frame(edit_frame)
        right_frame.pack(side=tk.RIGHT, fill=tk.Y, padx=5)
        
        ttk.Button(right_frame, text="新增题目", command=self._add_question, width=20).pack(pady=5)
        ttk.Button(right_frame, text="保存修改", command=self._update_question, width=20).pack(pady=5)
        ttk.Button(right_frame, text="删除题目", command=self._delete_question, width=20).pack(pady=5)
        ttk.Button(right_frame, text="批量删除", command=self._batch_delete_questions, width=20).pack(pady=5)
        ttk.Button(right_frame, text="清空编辑", command=self._clear_edit, width=20).pack(pady=5)
        ttk.Button(right_frame, text="AI提示词", command=self._show_ai_prompt, width=20).pack(pady=5)
        ttk.Button(right_frame, text="上传Word文档", command=self._upload_word, width=20).pack(pady=5)
    
    def _init_paper_page(self):
        """
        初始化组卷导出页
        """
        # 顶部参数设置区域
        param_frame = ttk.Frame(self.paper_frame)
        param_frame.pack(fill=tk.X, pady=5)
        
        # 题库选择
        ttk.Label(param_frame, text="题库:").pack(side=tk.LEFT, padx=5)
        self.paper_bank_var = tk.StringVar()
        self.paper_bank_combobox = ttk.Combobox(param_frame, textvariable=self.paper_bank_var, width=15)
        self.paper_bank_combobox.pack(side=tk.LEFT, padx=5)
        self.paper_bank_combobox.bind("<<ComboboxSelected>>", self._filter_paper_by_bank)
        
        # 知识点选择
        ttk.Label(param_frame, text="知识点:").pack(side=tk.LEFT, padx=5)
        # 创建知识点多选区域
        topic_frame = ttk.Frame(param_frame)
        topic_frame.pack(side=tk.LEFT, padx=5)
        
        # 知识点列表
        self.topic_listbox = tk.Listbox(topic_frame, selectmode=tk.MULTIPLE, width=30, height=5)
        self.topic_listbox.pack(side=tk.LEFT, padx=5)
        
        # 滚动条
        topic_scrollbar = ttk.Scrollbar(topic_frame, orient=tk.VERTICAL, command=self.topic_listbox.yview)
        topic_scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        self.topic_listbox.config(yscrollcommand=topic_scrollbar.set)
        
        ttk.Label(param_frame, text="题目数量:").pack(side=tk.LEFT, padx=5)
        self.question_count_var = tk.StringVar(value="10")
        ttk.Entry(param_frame, textvariable=self.question_count_var, width=10).pack(side=tk.LEFT, padx=5)
        
        ttk.Button(param_frame, text="生成试卷", command=self._generate_paper).pack(side=tk.LEFT, padx=5)
        
        # 中间显示区域
        display_frame = ttk.Frame(self.paper_frame)
        display_frame.pack(fill=tk.BOTH, expand=True, pady=5)
        
        # 左侧试卷显示
        paper_frame = ttk.LabelFrame(display_frame, text="试卷")
        paper_frame.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=5)
        
        self.paper_text = tk.Text(paper_frame, height=20, wrap=tk.WORD)
        self.paper_text.pack(fill=tk.BOTH, expand=True, pady=5)
        
        # 右侧答案显示
        answer_frame = ttk.LabelFrame(display_frame, text="答案")
        answer_frame.pack(side=tk.RIGHT, fill=tk.BOTH, expand=True, padx=5)
        
        self.answer_text = tk.Text(answer_frame, height=20, wrap=tk.WORD)
        self.answer_text.pack(fill=tk.BOTH, expand=True, pady=5)
        
        # 底部导出按钮区域
        export_frame = ttk.Frame(self.paper_frame)
        export_frame.pack(fill=tk.X, pady=5)
        
        ttk.Button(export_frame, text="导出试卷为Word", command=self._export_paper_word).pack(side=tk.LEFT, padx=5)
        ttk.Button(export_frame, text="导出试卷为PDF", command=self._export_paper_pdf).pack(side=tk.LEFT, padx=5)
        ttk.Button(export_frame, text="导出答案为Word", command=self._export_answers_word).pack(side=tk.LEFT, padx=5)
        ttk.Button(export_frame, text="导出答案为PDF", command=self._export_answers_pdf).pack(side=tk.LEFT, padx=5)
    
    def _load_banks(self):
        """
        加载题库列表到下拉框
        """
        banks = self.question_manager.get_banks()
        # 按名称排序题库
        sorted_banks = sorted(banks, key=lambda x: x['name'])
        bank_names = [bank['name'] for bank in sorted_banks]
        
        # 更新题目管理页的题库下拉框
        self.bank_combobox['values'] = ['所有题库'] + bank_names
        self.bank_combobox.current(0)  # 默认选择所有题库
        
        # 更新组卷导出页的题库下拉框
        if hasattr(self, 'paper_bank_combobox'):
            self.paper_bank_combobox['values'] = bank_names
            if bank_names:
                self.paper_bank_combobox.current(0)  # 默认选择第一个题库
                # 手动触发事件，更新知识点列表
                self._filter_paper_by_bank()
    
    def _load_questions(self, bank_id=None):
        """
        加载题目数据到列表
        
        Args:
            bank_id: 题库ID，None表示所有题库
        """
        # 清空列表
        for item in self.question_tree.get_children():
            self.question_tree.delete(item)
        
        # 加载所有题库信息
        banks = self.question_manager.get_banks()
        bank_map = {bank['id']: bank['name'] for bank in banks}
        
        # 加载题目
        if bank_id:
            # 获取指定题库的题目
            questions = self.question_manager.get_all_questions(bank_id)
            bank_name = bank_map.get(bank_id, '')
            for q in questions:
                self.question_tree.insert("", tk.END, values=(q['id'], bank_name, q['topic'], q['content']))
        else:
            # 获取所有题库的题目
            for bank in banks:
                bank_id = bank['id']
                bank_name = bank['name']
                questions = self.question_manager.get_all_questions(bank_id)
                for q in questions:
                    self.question_tree.insert("", tk.END, values=(q['id'], bank_name, q['topic'], q['content']))
        
        # 更新知识点下拉框
        topics = self.question_manager.get_all_topics(bank_id)
        self.topic_combobox['values'] = [''] + topics
        
        # 更新知识点多选列表
        self.topic_listbox.delete(0, tk.END)
        for topic in topics:
            self.topic_listbox.insert(tk.END, topic)
    
    def _filter_by_bank(self, event=None):
        """
        根据题库筛选题目
        """
        selected_bank = self.bank_var.get()
        
        if selected_bank == '所有题库':
            # 加载所有题库的题目
            self._load_questions()
        else:
            # 找到选中题库的ID
            banks = self.question_manager.get_banks()
            bank_id = None
            for bank in banks:
                if bank['name'] == selected_bank:
                    bank_id = bank['id']
                    break
            
            # 加载指定题库的题目
            if bank_id:
                self._load_questions(bank_id)
    
    def _filter_paper_by_bank(self, event=None):
        """
        在组卷导出页面根据题库更新知识点列表
        """
        selected_bank = self.paper_bank_var.get()
        
        if not selected_bank:
            return
        
        # 找到选中题库的ID
        banks = self.question_manager.get_banks()
        bank_id = None
        for bank in banks:
            if bank['name'] == selected_bank:
                bank_id = bank['id']
                break
        
        # 更新知识点多选列表
        if bank_id:
            topics = self.question_manager.get_all_topics(bank_id)
            self.topic_listbox.delete(0, tk.END)
            for topic in topics:
                self.topic_listbox.insert(tk.END, topic)
    
    def _filter_questions(self, event=None):
        """
        根据知识点筛选题目
        """
        topic = self.topic_var.get()
        selected_bank = self.bank_var.get()
        
        # 找到选中题库的ID
        bank_id = None
        if selected_bank != '所有题库':
            banks = self.question_manager.get_banks()
            for bank in banks:
                if bank['name'] == selected_bank:
                    bank_id = bank['id']
                    break
        
        # 清空列表
        for item in self.question_tree.get_children():
            self.question_tree.delete(item)
        
        # 筛选题目
        questions = self.question_manager.get_questions_by_topic(topic, bank_id)
        for q in questions:
            self.question_tree.insert("", tk.END, values=(q['id'], q['topic'], q['content']))
    
    def _search_questions(self):
        """
        根据关键词搜索题目
        """
        keyword = self.search_var.get()
        selected_bank = self.bank_var.get()
        
        # 找到选中题库的ID
        bank_id = None
        if selected_bank != '所有题库':
            banks = self.question_manager.get_banks()
            for bank in banks:
                if bank['name'] == selected_bank:
                    bank_id = bank['id']
                    break
        
        # 清空列表
        for item in self.question_tree.get_children():
            self.question_tree.delete(item)
        
        # 搜索题目
        questions = self.question_manager.search_questions(keyword, bank_id)
        for q in questions:
            self.question_tree.insert("", tk.END, values=(q['id'], q['topic'], q['content']))
    
    def _clear_search(self):
        """
        清空搜索
        """
        self.search_var.set("")
        self._load_questions()
    
    def _manage_banks(self):
        """
        题库管理窗口
        """
        # 创建新窗口
        bank_window = tk.Toplevel(self.root)
        bank_window.title("题库管理")
        bank_window.geometry("400x300")
        bank_window.resizable(True, True)
        
        # 添加列表区域
        list_frame = ttk.Frame(bank_window)
        list_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)
        
        # 创建滚动条
        scrollbar = ttk.Scrollbar(list_frame)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        
        # 创建题库列表
        bank_listbox = tk.Listbox(list_frame, selectmode=tk.SINGLE, yscrollcommand=scrollbar.set, width=40)
        bank_listbox.pack(fill=tk.BOTH, expand=True)
        scrollbar.config(command=bank_listbox.yview)
        
        # 加载题库列表
        def load_bank_list():
            bank_listbox.delete(0, tk.END)
            banks = self.question_manager.get_banks()
            # 按名称排序题库
            sorted_banks = sorted(banks, key=lambda x: x['name'])
            for bank in sorted_banks:
                bank_listbox.insert(tk.END, f"{bank['id']}. {bank['name']}")
        
        load_bank_list()
        
        # 添加操作区域
        action_frame = ttk.Frame(bank_window)
        action_frame.pack(fill=tk.X, padx=10, pady=10)
        
        # 添加题库
        def add_bank():
            from tkinter import simpledialog
            bank_name = simpledialog.askstring("添加题库", "请输入题库名称:", parent=bank_window)
            if bank_name:
                self.question_manager.add_bank(bank_name)
                load_bank_list()
                self._load_banks()  # 更新主窗口的题库下拉框
        
        # 修改题库名称
        def rename_bank():
            selected_index = bank_listbox.curselection()
            if not selected_index:
                messagebox.showinfo("提示", "请先选择要修改的题库")
                return
            
            selected_text = bank_listbox.get(selected_index[0])
            bank_id = int(selected_text.split('.')[0])
            
            from tkinter import simpledialog
            new_name = simpledialog.askstring("修改题库名称", "请输入新的题库名称:", parent=bank_window)
            if new_name:
                self.question_manager.update_bank_name(bank_id, new_name)
                load_bank_list()
                self._load_banks()  # 更新主窗口的题库下拉框
        
        # 删除题库
        def delete_bank():
            selected_index = bank_listbox.curselection()
            if not selected_index:
                messagebox.showinfo("提示", "请先选择要删除的题库")
                return
            
            selected_text = bank_listbox.get(selected_index[0])
            bank_id = int(selected_text.split('.')[0])
            
            # 检查是否是最后一个题库
            banks = self.question_manager.get_banks()
            if len(banks) <= 1:
                messagebox.showinfo("提示", "至少需要保留一个题库")
                return
            
            if messagebox.askyesno("确认", "确定要删除这个题库吗？"):
                self.question_manager.delete_bank(bank_id)
                load_bank_list()
                self._load_banks()  # 更新主窗口的题库下拉框
                self._load_questions()  # 重新加载题目列表
        
        # 添加按钮
        ttk.Button(action_frame, text="添加题库", command=add_bank, width=15).pack(side=tk.LEFT, padx=5)
        ttk.Button(action_frame, text="修改名称", command=rename_bank, width=15).pack(side=tk.LEFT, padx=5)
        ttk.Button(action_frame, text="删除题库", command=delete_bank, width=15).pack(side=tk.LEFT, padx=5)
        ttk.Button(action_frame, text="关闭", command=bank_window.destroy, width=15).pack(side=tk.RIGHT, padx=5)
    
    def _on_question_select(self, event):
        """
        当选择题目时的处理
        """
        selected_items = self.question_tree.selection()
        if not selected_items:
            return
        
        # 只处理第一个选中的题目，即使是多选模式
        item = selected_items[0]
        values = self.question_tree.item(item, "values")
        question_id = int(values[0])
        
        # 获取题目详情
        question, bank_id = self.question_manager.get_question(question_id)
        if question:
            self.current_question_id = question_id
            self.edit_topic_var.set(question['topic'])
            self.edit_content_text.delete(1.0, tk.END)
            self.edit_content_text.insert(tk.END, question['content'])
            self.edit_answer_text.delete(1.0, tk.END)
            self.edit_answer_text.insert(tk.END, question['answer'])
    
    def _add_question(self):
        """
        新增题目
        """
        topic = self.edit_topic_var.get()
        content = self.edit_content_text.get(1.0, tk.END).strip()
        answer = self.edit_answer_text.get(1.0, tk.END).strip()
        
        if not topic or not content or not answer:
            messagebox.showerror("错误", "请填写完整的题目信息")
            return
        
        try:
            # 获取当前选中的题库ID
            selected_bank = self.bank_var.get()
            bank_id = None
            
            if selected_bank == '所有题库':
                # 如果选择了所有题库，默认使用第一个题库
                banks = self.question_manager.get_banks()
                if banks:
                    bank_id = banks[0]['id']
            else:
                # 找到选中题库的ID
                banks = self.question_manager.get_banks()
                for bank in banks:
                    if bank['name'] == selected_bank:
                        bank_id = bank['id']
                        break
            
            if not bank_id:
                messagebox.showerror("错误", "未找到可用的题库")
                return
            
            question_data = {
                'topic': topic,
                'content': content,
                'answer': answer
            }
            self.question_manager.add_question(question_data, bank_id)
            
            # 重新加载题目列表
            selected_bank = self.bank_var.get()
            if selected_bank != '所有题库':
                # 找到选中题库的ID
                banks = self.question_manager.get_banks()
                bank_id = None
                for bank in banks:
                    if bank['name'] == selected_bank:
                        bank_id = bank['id']
                        break
                self._load_questions(bank_id)
            else:
                self._load_questions()
            
            # 显示自动消失的成功消息
            success_window = tk.Toplevel(self.root)
            success_window.title("成功")
            success_window.geometry("300x100")
            success_window.transient(self.root)  # 设置为临时窗口
            
            # 设置窗口位置在主窗口中央
            x = self.root.winfo_x() + (self.root.winfo_width() // 2) - 150
            y = self.root.winfo_y() + (self.root.winfo_height() // 2) - 50
            success_window.geometry(f"300x100+{x}+{y}")
            
            # 添加成功消息
            ttk.Label(success_window, text="题目添加成功", font=("SimHei", 12)).pack(pady=20)
            
            # 1秒后自动关闭窗口
            success_window.after(1000, success_window.destroy)
            
            if self.continuous_add_mode:
                # 如果已经在连续添加模式下，直接清空题目内容和答案
                self.edit_content_text.delete(1.0, tk.END)
                self.edit_answer_text.delete(1.0, tk.END)
            else:
                # 询问用户是否需要重复导入同一知识点的题目
                repeat = messagebox.askyesno("重复导入", "是否继续添加同一知识点的题目？")
                
                if repeat:
                    # 进入连续添加模式
                    self.continuous_add_mode = True
                    # 保持知识点不变，只清空题目内容和答案
                    self.edit_content_text.delete(1.0, tk.END)
                    self.edit_answer_text.delete(1.0, tk.END)
                    messagebox.showinfo("提示", "已进入连续添加模式，将保持知识点不变")
                else:
                    # 清空所有输入
                    self._clear_edit()
        except Exception as e:
            messagebox.showerror("错误", f"添加题目失败: {e}")
    
    def _update_question(self):
        """
        更新题目
        """
        if not self.current_question_id:
            messagebox.showerror("错误", "请先选择要修改的题目")
            return
        
        topic = self.edit_topic_var.get()
        content = self.edit_content_text.get(1.0, tk.END).strip()
        answer = self.edit_answer_text.get(1.0, tk.END).strip()
        
        if not topic or not content or not answer:
            messagebox.showerror("错误", "请填写完整的题目信息")
            return
        
        try:
            question_data = {
                'topic': topic,
                'content': content,
                'answer': answer
            }
            self.question_manager.update_question(self.current_question_id, question_data)
            self._load_questions()
            messagebox.showinfo("成功", "题目修改成功")
        except Exception as e:
            messagebox.showerror("错误", f"修改题目失败: {e}")
    
    def _delete_question(self):
        """
        删除题目
        """
        if not self.current_question_id:
            messagebox.showerror("错误", "请先选择要删除的题目")
            return
        
        if messagebox.askyesno("确认", "确定要删除这个题目吗？"):
            try:
                self.question_manager.delete_question(self.current_question_id)
                self._load_questions()
                self._clear_edit()
                messagebox.showinfo("成功", "题目删除成功")
            except Exception as e:
                messagebox.showerror("错误", f"删除题目失败: {e}")
    
    def _batch_delete_questions(self):
        """
        批量删除题目
        """
        # 获取选中的题目
        selected_items = self.question_tree.selection()
        
        if not selected_items:
            messagebox.showerror("错误", "请先选择要删除的题目")
            return
        
        # 提取选中题目的ID
        selected_ids = []
        for item in selected_items:
            values = self.question_tree.item(item, "values")
            if values:
                selected_ids.append(int(values[0]))
        
        # 显示确认对话框
        if messagebox.askyesno("确认", f"确定要删除这 {len(selected_ids)} 道题目吗？"):
            try:
                # 批量删除题目
                deleted_count = 0
                for question_id in selected_ids:
                    self.question_manager.delete_question(question_id)
                    deleted_count += 1
                
                # 重新加载题目列表
                self._load_questions()
                
                # 清空编辑区域
                self._clear_edit()
                
                # 显示删除结果
                messagebox.showinfo("成功", f"成功删除 {deleted_count} 道题目")
            except Exception as e:
                messagebox.showerror("错误", f"删除题目失败: {e}")
    
    def _show_ai_prompt(self):
        """
        显示AI提示词，即题目格式要求
        """
        # 创建新窗口
        prompt_window = tk.Toplevel(self.root)
        prompt_window.title("AI提示词 - 题目格式要求")
        prompt_window.geometry("500x300")
        prompt_window.resizable(True, True)
        
        # 添加文本区域
        text_frame = ttk.Frame(prompt_window)
        text_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)
        
        # 创建滚动条
        scrollbar = ttk.Scrollbar(text_frame)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        
        # 创建文本框，设置为只读模式
        text_widget = tk.Text(text_frame, wrap=tk.WORD, yscrollcommand=scrollbar.set, font=("SimHei", 10))
        text_widget.pack(fill=tk.BOTH, expand=True)
        scrollbar.config(command=text_widget.yview)
        
        # 插入提示文本
        prompt_text = "我现在要确保我的题目的格式是：\n"
        prompt_text += "- 按照指定格式编写题目，确保每个题目以题号开头（如\"1. \"）\n"
        prompt_text += "- 选项以\"A. \"、\"B. \"等格式开头\n"
        prompt_text += "- 答案单独一行，只包含选项字母（如\"A\"）\n"
        prompt_text += "- 删除带有图片的题目\n"
        prompt_text += "\n请严格按照格式要求进行修改之后，以HTML形式发送给我"
        
        text_widget.insert(tk.END, prompt_text)
        text_widget.config(state=tk.DISABLED)  # 设置为只读
        
        # 添加关闭按钮
        button_frame = ttk.Frame(prompt_window)
        button_frame.pack(fill=tk.X, padx=10, pady=5)
        
        ttk.Button(button_frame, text="关闭", command=prompt_window.destroy, width=15).pack(side=tk.RIGHT, padx=5, pady=5)
    
    def _clear_edit(self):
        """
        清空编辑区域
        """
        self.current_question_id = None
        self.edit_topic_var.set("")
        self.edit_content_text.delete(1.0, tk.END)
        self.edit_answer_text.delete(1.0, tk.END)
    
    def _exit_topic(self):
        """
        退出当前知识点输入，清空知识点输入框
        """
        # 退出连续添加模式
        self.continuous_add_mode = False
        # 清空知识点输入框
        self.edit_topic_var.set("")
        # 重新加载所有知识点
        all_topics = self.question_manager.get_all_topics()
        self.topic_combobox_edit['values'] = all_topics
        # 清空题目内容和答案
        self.edit_content_text.delete(1.0, tk.END)
        self.edit_answer_text.delete(1.0, tk.END)
        messagebox.showinfo("提示", "已退出连续添加模式，可输入新的知识点")
    
    def _sort_column(self, col):
        """
        排序列
        
        Args:
            col: 列名
        """
        # 获取所有项目并转换为列表
        items = list(self.question_tree.get_children(''))
        # 提取项目值并排序
        items.sort(key=lambda x: int(self.question_tree.set(x, col)))
        
        # 重新插入排序后的项目
        for item in items:
            self.question_tree.move(item, '', tk.END)
    
    def _upload_word(self):
        """
        上传Word文档并解析题目
        """
        try:
            # 导入必要的模块
            from tkinter import filedialog
            try:
                from docx import Document
            except ImportError:
                messagebox.showerror("错误", "缺少python-docx库，请先安装依赖")
                return
            
            # 打开文件选择对话框
            file_path = filedialog.askopenfilename(
                title="选择Word文档",
                filetypes=[("Word文档", "*.docx"), ("所有文件", "*.*")]
            )
            
            if not file_path:
                return  # 用户取消选择
            
            # 解析Word文档
            questions = self._parse_word_document(file_path)
            
            if not questions:
                messagebox.showinfo("提示", "文档中未找到题目")
                return
            
            # 询问用户为解析出的题目设置知识点
            from tkinter import simpledialog
            topic = simpledialog.askstring(
                "设置知识点",
                f"已解析出 {len(questions)} 道题目，请输入知识点:",
                parent=self.root
            )
            
            if not topic:
                return  # 用户取消输入
            
            # 获取当前选中的题库ID
            selected_bank = self.bank_var.get()
            bank_id = None
            
            if selected_bank == '所有题库':
                # 如果选择了所有题库，默认使用第一个题库
                banks = self.question_manager.get_banks()
                if banks:
                    bank_id = banks[0]['id']
            else:
                # 找到选中题库的ID
                banks = self.question_manager.get_banks()
                for bank in banks:
                    if bank['name'] == selected_bank:
                        bank_id = bank['id']
                        break
            
            if not bank_id:
                messagebox.showerror("错误", "未找到可用的题库")
                return
            
            # 批量添加题目
            added_count = 0
            for q in questions:
                try:
                    question_data = {
                        'topic': topic,
                        'content': q['content'],
                        'answer': q['answer']
                    }
                    self.question_manager.add_question(question_data, bank_id)
                    added_count += 1
                except Exception as e:
                    messagebox.showerror("错误", f"添加题目失败: {e}")
                    continue
            
            # 重新加载题目列表
            self._load_questions()
            
            # 显示添加结果
            messagebox.showinfo("成功", f"成功添加 {added_count} 道题目到知识点 '{topic}'")
            
        except Exception as e:
            messagebox.showerror("错误", f"上传失败: {e}")
    
    def _parse_word_document(self, file_path):
        """
        解析Word文档中的题目
        
        Args:
            file_path: Word文档路径
            
        Returns:
            list: 解析出的题目列表
        """
        from docx import Document
        
        doc = Document(file_path)
        questions = []
        
        current_question_content = []
        current_answer = None
        in_question = False
        
        for para in doc.paragraphs:
            text = para.text.strip()
            
            # 检查是否是新题目开始
            if text.startswith(tuple(str(i) + '.' for i in range(1, 1000))):
                # 如果有未完成的题目，保存它
                if in_question and current_question_content and current_answer:
                    questions.append({
                        'content': ' '.join(current_question_content), 
                        'answer': current_answer
                    })
                
                # 开始新题目
                current_question_content = [text]
                current_answer = None
                in_question = True
            
            # 检查是否是答案行（单个选项字母或多个选项组合）
            elif in_question and text and all(c in 'ABCD' for c in text):
                current_answer = text
            
            # 检查是否是题目内容的一部分（包括选项）
            elif in_question and text:
                current_question_content.append(text)
        
        # 保存最后一道题目
        if in_question and current_question_content and current_answer:
            questions.append({
                'content': ' '.join(current_question_content), 
                'answer': current_answer
            })
        
        # 清理题目内容和答案
        for q in questions:
            # 清理题目内容，去除题号
            content = q['content']
            # 找到第一个小数点的位置
            dot_pos = content.find('.')
            if dot_pos != -1:
                # 去除题号，保留剩余内容
                q['content'] = content[dot_pos + 1:].strip()
            
            # 确保答案格式正确
            answer = q['answer'].strip()
            if answer in ('A', 'B', 'C', 'D'):
                q['answer'] = answer
            elif answer.startswith(('A.', 'B.', 'C.', 'D.')):
                q['answer'] = answer[0]
        
        return questions
    
    def _on_topic_input(self, event=None):
        """
        知识点输入时的自动匹配功能
        
        Args:
            event: 键盘事件
        """
        # 获取当前输入的文本
        current_input = self.edit_topic_var.get().lower()
        
        # 获取所有知识点
        all_topics = self.question_manager.get_all_topics()
        
        # 筛选匹配的知识点
        if current_input:
            matched_topics = [topic for topic in all_topics if current_input in topic.lower()]
        else:
            matched_topics = all_topics
        
        # 更新Combobox的下拉列表
        self.topic_combobox_edit['values'] = matched_topics
        
        # 确保Combobox为可编辑模式，允许输入新知识点
        self.topic_combobox_edit['state'] = 'normal'
        
        # 只在用户按下特定键时才显示下拉列表，避免干扰用户输入
        # 这里不自动显示下拉列表，让用户通过点击下拉箭头或按向下键来查看匹配结果
    
    def _generate_paper(self):
        """
        生成试卷
        """
        # 获取选中的题库
        selected_bank = self.paper_bank_var.get()
        if not selected_bank:
            messagebox.showerror("错误", "请选择一个题库")
            return
        
        # 找到选中题库的ID
        banks = self.question_manager.get_banks()
        selected_bank_id = None
        for bank in banks:
            if bank['name'] == selected_bank:
                selected_bank_id = bank['id']
                break
        
        if not selected_bank_id:
            messagebox.showerror("错误", "未找到选中的题库")
            return
        
        # 获取选中的知识点
        selected_indices = self.topic_listbox.curselection()
        selected_topics = [self.topic_listbox.get(i) for i in selected_indices]
        
        # 如果没有选择知识点，使用所有知识点
        if not selected_topics:
            selected_topics = None
        
        try:
            question_count = int(self.question_count_var.get())
        except ValueError:
            messagebox.showerror("错误", "请输入有效的题目数量")
            return
        
        try:
            # 获取指定题库的题目
            questions = self.question_manager.get_all_questions(selected_bank_id)
            
            # 根据知识点筛选题目
            if selected_topics:
                filtered_questions = []
                for q in questions:
                    if q.get('topic') in selected_topics:
                        filtered_questions.append(q)
                questions = filtered_questions
            
            if not questions:
                messagebox.showinfo("提示", "未找到符合条件的题目")
                return
            
            # 随机打乱题目顺序
            import random
            random.shuffle(questions)
            
            # 截取指定数量的题目
            actual_count = min(question_count, len(questions))
            selected_questions = questions[:actual_count]
            
            # 生成试卷题目列表（不含答案）
            paper_questions = []
            for i, q in enumerate(selected_questions, 1):
                paper_question = {
                    'id': i,
                    'content': q['content'],
                    'topic': q['topic']
                }
                paper_questions.append(paper_question)
            
            # 生成答案列表
            answers = []
            for i, q in enumerate(selected_questions, 1):
                answer = {
                    'id': i,
                    'content': q['content'],
                    'answer': q['answer'],
                    'topic': q['topic']
                }
                answers.append(answer)
            
            self.current_paper = paper_questions
            self.current_answers = answers
            
            # 显示试卷
            self.paper_text.delete(1.0, tk.END)
            for q in paper_questions:
                self.paper_text.insert(tk.END, f"{q['id']}. {q['topic']}\n")
                self.paper_text.insert(tk.END, f"{q['content']}\n\n")
            
            # 显示答案
            self.answer_text.delete(1.0, tk.END)
            for ans in answers:
                self.answer_text.insert(tk.END, f"{ans['id']}. {ans['topic']}\n")
                self.answer_text.insert(tk.END, f"题目: {ans['content']}\n")
                self.answer_text.insert(tk.END, f"答案: {ans['answer']}\n\n")
            
            messagebox.showinfo("成功", "试卷生成成功")
        except Exception as e:
            messagebox.showerror("错误", f"生成试卷失败: {e}")
    
    def _export_paper_word(self):
        """
        导出试卷为Word
        """
        if not self.current_paper:
            messagebox.showerror("错误", "请先生成试卷")
            return
        
        try:
            self.exporter.export_paper_to_word(self.current_paper)
            messagebox.showinfo("成功", "试卷导出为Word成功")
        except Exception as e:
            messagebox.showerror("错误", f"导出失败: {e}")
    
    def _export_paper_pdf(self):
        """
        导出试卷为PDF
        """
        if not self.current_paper:
            messagebox.showerror("错误", "请先生成试卷")
            return
        
        try:
            self.exporter.export_paper_to_pdf(self.current_paper)
            messagebox.showinfo("成功", "试卷导出为PDF成功")
        except Exception as e:
            messagebox.showerror("错误", f"导出失败: {e}")
    
    def _export_answers_word(self):
        """
        导出答案为Word
        """
        if not self.current_answers:
            messagebox.showerror("错误", "请先生成试卷")
            return
        
        try:
            self.exporter.export_answers_to_word(self.current_answers)
            messagebox.showinfo("成功", "答案导出为Word成功")
        except Exception as e:
            messagebox.showerror("错误", f"导出失败: {e}")
    
    def _export_answers_pdf(self):
        """
        导出答案为PDF
        """
        if not self.current_answers:
            messagebox.showerror("错误", "请先生成试卷")
            return
        
        try:
            self.exporter.export_answers_to_pdf(self.current_answers)
            messagebox.showinfo("成功", "答案导出为PDF成功")
        except Exception as e:
            messagebox.showerror("错误", f"导出失败: {e}")
